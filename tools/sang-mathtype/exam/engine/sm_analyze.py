"""Các hàm dùng chung cho cả bản máy (examd.py) và bản trình duyệt (Pyodide trong add-in).
Chỉ dùng thư viện chuẩn + docx_shuffler/openpyxl; hàm nào có await thì chạy được trong vòng lặp sự kiện của Pyodide."""
import io, json, inspect, zipfile


class _Py:
    """Bọc dữ liệu để đưa cho hàm cũ vốn mong đợi proxy JS có .to_py()."""
    def __init__(self, value): self.value = value
    def to_py(self): return self.value


async def _maybe(value):
    return await value if inspect.isawaitable(value) else value


async def analyze_exam(data):
    """Đếm câu từng phần và nhận đáp án đã đánh dấu sẵn trong đề (gạch chân chữ cái, Đ/S, đáp án ngắn).
    Dùng đúng bộ trộn đề nên nhận diện khớp với kết quả trộn: chạy một lượt giữ nguyên thứ tự rồi đọc lại đáp án."""
    import docx_shuffler
    count = json.loads(await _maybe(docx_shuffler.count_questions(_Py(data))))
    if count.get('error'):
        return {'count': count, 'answers': None}
    try:
        _stream, a1, a2, a3 = docx_shuffler.shuffle_questions_in_docx(io.BytesIO(data), {'ma_de': '000'}, True)
        return {'count': count, 'answers': {'mc': [(a[0] if a else '') for a in a1], 'tf': list(a2), 'sa': [str(x) for x in a3]}}
    except Exception as e:  # đếm được nhưng không đọc được đáp án: báo rõ thay vì im lặng
        return {'count': count, 'answers': None, 'answer_error': f'{type(e).__name__}: {e}'}


async def analyze_exam_js(file_content_js):
    return json.dumps(await analyze_exam(file_content_js.to_py()), ensure_ascii=False)


async def shuffle_pack(file_content_js, options_js):
    """Trộn đề và trả về {files: {tên: bytes}, key: [[mã đề, đáp án câu 1, ...], ...]} để trình duyệt tự mở từng mã đề."""
    import docx_shuffler
    data = file_content_js.to_py()
    options = options_js.to_py()
    zip_bytes = await _maybe(docx_shuffler.process_file(_Py(data), _Py(options)))
    files = {}
    with zipfile.ZipFile(io.BytesIO(bytes(zip_bytes))) as z:
        for info in z.infolist():
            if not info.is_dir():
                files[info.filename.split('/')[-1]] = z.read(info)
    key = []
    if 'ChamThi.xlsx' in files:
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(files['ChamThi.xlsx'])).active
        key = [[('' if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
    return {'files': files, 'key': key}
